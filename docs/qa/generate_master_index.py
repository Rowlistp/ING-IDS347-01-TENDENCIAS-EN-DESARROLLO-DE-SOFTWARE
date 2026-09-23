# -*- coding: utf-8 -*-
"""
Generador de Índice Maestro Consolidado de Defectos (.DOCX)
Sistema de Gestión y Control de Combustible — FuelTrack ERP
Aseguramiento de Calidad (QA) — Rowlis Trinidad <Rowlistp@gmail.com>
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

COLOR_TANQUE = "16333A"
COLOR_ACERO = "4A5A63"
COLOR_MEDIDOR = "E29B2E"
COLOR_EXITO = "2E7D5B"
COLOR_PELIGRO = "C1432B"
COLOR_FONDO = "F7F8F6"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, hex_color="CCCCCC"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

ITEMS = [
    ("DEF-2026-001", "Seguridad & API", "Expiración de Token JWT y Redirección Insegura", "Crítica", "P1", "Cerrado", "DEF-2026-001_Vulnerabilidad_Sesion_JWT_Expirada_Y_Redireccion.docx"),
    ("DEF-2026-002", "Cierre Diario", "Generación On-Demand y Persistencia de Acta PDF", "Mayor", "P1", "Cerrado", "DEF-2026-002_Ausencia_Generacion_OnDemand_PDF_Cierre_Diario.docx"),
    ("DEF-2026-003", "Despachos & Tickets", "Sellos Criptográficos, QR y 3 Firmas en Ticket PDF", "Mayor", "P1", "Cerrado", "DEF-2026-003_Falta_Sellos_Criptograficos_QR_Y_Firmas_Ticket_PDF.docx"),
    ("DEF-2026-004", "Validaciones & Entradas", "Validaciones Dominicanas (RNC, Cédula, Teléfonos 809)", "Crítica", "P1", "Cerrado", "DEF-2026-004_Inexistencia_Validaciones_Dominio_Republica_Dominicana.docx"),
    ("DEF-2026-005", "UI/UX & Ergonomía", "Contraste StatusBadge y Botones Táctiles (38px)", "Media", "P2", "Cerrado", "DEF-2026-005_Deficiencia_Contraste_StatusBadge_Y_Ergonomia_Tactil.docx"),
    ("DEF-2026-006", "Recepciones Combustible", "Verificación Espacio en Tanque y Comprobante Digital", "Mayor", "P1", "Cerrado", "DEF-2026-006_Falta_Validacion_Espacio_Tanque_Y_Comprobante_Recepciones.docx"),
    ("DEF-2026-007", "Tanques Estacionarios", "Asistente Nomenclatura y Regla Nivel Crítico Menor", "Media", "P2", "Cerrado", "DEF-2026-007_Ambiguedad_Nomenclatura_Tanques_Y_Nivel_Critico.docx"),
    ("DEF-2026-008", "Reportes & Notificaciones", "Bloqueo Fechas Invertidas (Desde > Hasta)", "Media", "P2", "Cerrado", "DEF-2026-008_Carencia_Control_Fechas_Invertidas_Reportes_Notificaciones.docx"),
    ("DEF-2026-009", "Solicitudes Combustible", "Cálculos Rápidos Aprobación y Cuota Autorizada", "Media", "P2", "Cerrado", "DEF-2026-009_Ausencia_Calculo_Rapido_Aprobacion_Solicitudes.docx"),
    ("DEF-2026-010", "Gestión de Tickets", "Selector Prefijos Oficiales y Desacoplamiento Tabla", "Mayor", "P1", "Cerrado", "DEF-2026-010_Falta_Nomenclatura_Controlada_Y_Auditoria_Tickets.docx"),
    ("DEF-2026-011", "Despachos & Estaciones", "Bloqueo por Estación Inactiva en Frontend", "Mayor", "P1", "Cerrado", "DEF-2026-011_Bloqueo_Despacho_Estacion_Inactiva_Frontend.docx"),
    ("DEF-2026-012", "Tanques & Inventario", "Bloqueo Creación Tanque con Combustible Inactivo", "Mayor", "P1", "Cerrado", "DEF-2026-012_Creacion_Tanque_Combustible_Inactivo_Permitida.docx"),
    ("DEF-2026-013", "Solicitudes Recurrentes", "Bloqueo Entidades Inactivas y Derivación Depto", "Mayor", "P1", "Cerrado", "DEF-2026-013_Solicitudes_Recurrentes_Entidades_Inactivas_Y_Departamento.docx"),
    ("DEF-2026-014", "Solicitudes Combustible", "Validación y Derivación de Departamento de Empleado", "Mayor", "P1", "Cerrado", "DEF-2026-014_Solicitudes_Departamento_Inconsistente_Con_Empleado.docx"),
    ("DEF-2026-015", "Recepciones Combustible", "Bloqueo Recepción en Tanque con Combustible Inactivo", "Mayor", "P1", "Cerrado", "DEF-2026-015_Recepcion_Tanque_Combustible_Inactivo_Permitida.docx"),
    ("DEF-2026-016", "Inventario & Tanques", "Bloqueo Ajuste y Transferencia con Combustible Inactivo", "Mayor", "P1", "Cerrado", "DEF-2026-016_Inventario_Ajuste_Y_Transferencia_Combustible_Inactivo.docx"),
]

def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("SISTEMA FUELTRACK ERP  |  ÍNDICE MAESTRO DE CALIDAD  |  DOC-QA-00")
        hr.font.name = "Arial"
        hr.font.size = Pt(8)
        hr.font.color.rgb = RGBColor(0x4A, 0x5A, 0x63)
        
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = fp.add_run("DOCUMENTO TÉCNICO OFICIAL — CONSOLIDADO MAESTRO DE NO CONFORMIDADES  •  PÁGINA 1")
        fr.font.name = "Arial"
        fr.font.size = Pt(8)
        fr.font.color.rgb = RGBColor(0x4A, 0x5A, 0x63)

    # Banner Principal
    t_banner = doc.add_table(rows=1, cols=1)
    t_banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_banner = t_banner.rows[0].cells[0]
    set_cell_background(c_banner, COLOR_TANQUE)
    set_cell_margins(c_banner, top=160, bottom=160, left=200, right=200)
    p_b = c_banner.paragraphs[0]
    p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_b1 = p_b.add_run("ÍNDICE MAESTRO DE NO CONFORMIDADES Y CONTROL DE CALIDAD\n")
    r_b1.bold = True
    r_b1.font.name = "Arial"
    r_b1.font.size = Pt(14)
    r_b1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    r_b2 = p_b.add_run("SISTEMA DE GESTIÓN Y CONTROL DE COMBUSTIBLE — FUELTRACK ERP\n")
    r_b2.bold = True
    r_b2.font.name = "Arial"
    r_b2.font.size = Pt(10)
    r_b2.font.color.rgb = RGBColor(0xE2, 0x9B, 0x2E)

    r_b3 = p_b.add_run("TOTAL: 16 DEFECTOS AUDITADOS  •  16 CERRADOS (100% RESUELTOS)")
    r_b3.bold = True
    r_b3.font.name = "Arial"
    r_b3.font.size = Pt(9.5)
    r_b3.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    doc.add_paragraph()

    # Tabla Resumen Ejecutivo de Métricas
    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("1. RESUMEN EJECUTIVO DE ESTABILIZACIÓN Y CALIDAD")
    r_sub.bold = True
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(0x16, 0x33, 0x3A)

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.15
    p_intro.add_run(
        "El presente índice consolida formalmente los dieciséis (16) informes técnicos individuales de defectos, "
        "mejoras de interfaz, validaciones de dominio dominicano y robustecimiento de seguridad implementados "
        "en la plataforma FuelTrack. Todas las no conformidades fueron verificadas con 100% de pruebas exitosas "
        "en backend (326/326 pruebas unitarias en .NET 10) y compilación limpia sin errores de linter en frontend."
    ).font.size = Pt(9.5)

    t_res = doc.add_table(rows=5, cols=5)
    t_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_res, "B0C4DE")
    
    res_headers = ["Severidad", "Abiertos", "Resueltos", "Mitigados", "Total"]
    for i, h in enumerate(res_headers):
        c = t_res.cell(0, i)
        set_cell_background(c, COLOR_ACERO)
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        rh = c.paragraphs[0].add_run(h)
        rh.bold = True
        rh.font.name = "Arial"
        rh.font.size = Pt(9)
        rh.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    res_data = [
        ("Crítica", "0", "2", "0", "2"),
        ("Mayor", "0", "10", "0", "10"),
        ("Media", "0", "4", "0", "4"),
        ("Total General", "0", "16", "0", "16"),
    ]

    for r_idx, row in enumerate(res_data, start=1):
        for c_idx, val in enumerate(row):
            c = t_res.cell(r_idx, c_idx)
            set_cell_margins(c, top=60, bottom=60, left=100, right=100)
            if r_idx == 4:
                set_cell_background(c, COLOR_FONDO)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if r_idx == 4 or c_idx == 0:
                r.bold = True
            if c_idx == 2:
                r.font.color.rgb = RGBColor(0x2E, 0x7D, 0x5B)
                r.bold = True

    doc.add_paragraph()

    # Tabla Maestra de Defectos
    p_tab = doc.add_paragraph()
    r_tab = p_tab.add_run("2. RELACIÓN MAESTRA DE INFORMES TÉCNICOS AUDITADOS")
    r_tab.bold = True
    r_tab.font.name = "Arial"
    r_tab.font.size = Pt(11)
    r_tab.font.color.rgb = RGBColor(0x16, 0x33, 0x3A)

    t_master = doc.add_table(rows=len(ITEMS) + 1, cols=6)
    t_master.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_master, "CCCCCC")

    m_headers = ["Código", "Módulo", "Descripción de Corrección", "Severidad", "Prioridad", "Estado"]
    for i, h in enumerate(m_headers):
        c = t_master.cell(0, i)
        set_cell_background(c, COLOR_TANQUE)
        set_cell_margins(c, top=80, bottom=80, left=80, right=80)
        rh = c.paragraphs[0].add_run(h)
        rh.bold = True
        rh.font.name = "Arial"
        rh.font.size = Pt(8.5)
        rh.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, (cod, mod, desc, sev, prio, est, doc_name) in enumerate(ITEMS, start=1):
        c0 = t_master.cell(r_idx, 0)
        c1 = t_master.cell(r_idx, 1)
        c2 = t_master.cell(r_idx, 2)
        c3 = t_master.cell(r_idx, 3)
        c4 = t_master.cell(r_idx, 4)
        c5 = t_master.cell(r_idx, 5)

        for cell in [c0, c1, c2, c3, c4, c5]:
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)

        r0 = c0.paragraphs[0].add_run(cod)
        r0.bold = True
        r0.font.name = "Arial"
        r0.font.size = Pt(8)
        r0.font.color.rgb = RGBColor(0x16, 0x33, 0x3A)

        c1.paragraphs[0].add_run(mod).font.size = Pt(8)
        c2.paragraphs[0].add_run(desc).font.size = Pt(8)
        
        r3 = c3.paragraphs[0].add_run(sev)
        r3.font.size = Pt(8)
        if sev == "Crítica":
            r3.font.color.rgb = RGBColor(0xC1, 0x43, 0x2B)
            r3.bold = True
        elif sev == "Mayor":
            r3.font.color.rgb = RGBColor(0x85, 0x4D, 0x0E)

        c4.paragraphs[0].add_run(prio).font.size = Pt(8)

        r5 = c5.paragraphs[0].add_run("✅ " + est)
        r5.bold = True
        r5.font.size = Pt(8)
        r5.font.color.rgb = RGBColor(0x2E, 0x7D, 0x5B)

    doc.add_paragraph()

    # Firmas
    p_firmas = doc.add_paragraph()
    r_f_t = p_firmas.add_run("3. DICTAMEN FINAL Y HOMOLOGACIÓN DE ASEGURAMIENTO DE CALIDAD")
    r_f_t.bold = True
    r_f_t.font.name = "Arial"
    r_f_t.font.size = Pt(11)
    r_f_t.font.color.rgb = RGBColor(0x16, 0x33, 0x3A)

    t_f = doc.add_table(rows=1, cols=3)
    t_f.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_f, "FFFFFF")

    firmas = [
        ("Rowlis Trinidad", "Ingeniero Líder de QA", "Dictamen: Aprobado 100%"),
        ("Comité Técnico", "Arquitectura & Backend", "Homologado en main"),
        ("Dirección General", "Operaciones FuelTrack", "Certificación Válida")
    ]

    for i, (nom, cargo, dic) in enumerate(firmas):
        c = t_f.cell(0, i)
        set_cell_margins(c, top=140, bottom=60, left=60, right=60)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        p.add_run("_______________________________\n").font.color.rgb = RGBColor(0x4A, 0x5A, 0x63)
        
        r1 = p.add_run(nom + "\n")
        r1.bold = True
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = RGBColor(0x16, 0x33, 0x3A)
        
        r2 = p.add_run(cargo + "\n")
        r2.font.size = Pt(8)
        r2.font.color.rgb = RGBColor(0x4A, 0x5A, 0x63)
        
        r3 = p.add_run(dic)
        r3.bold = True
        r3.font.size = Pt(7.5)
        r3.font.color.rgb = RGBColor(0x2E, 0x7D, 0x5B)

    base_dir = os.path.dirname(os.path.abspath(__file__))
    out_file = os.path.join(base_dir, "00_INDICE_MAESTRO_INFORMES_DEFECTOS.docx")
    doc.save(out_file)
    print(f"[OK] Generado: 00_INDICE_MAESTRO_INFORMES_DEFECTOS.docx")

if __name__ == "__main__":
    main()
