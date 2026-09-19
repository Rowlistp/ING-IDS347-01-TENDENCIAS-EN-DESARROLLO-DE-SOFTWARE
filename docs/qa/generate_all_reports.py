# -*- coding: utf-8 -*-
"""
Generador Automatizado de Informes Técnicos Individuales de Defectos (.DOCX)
Sistema de Gestión y Control de Combustible — FuelTrack ERP
Aseguramiento de Calidad (QA) — Rowlis Trinidad <Rowlistp@gmail.com>
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

COLOR_TANQUE = "16333A"
COLOR_ACERO = "4A5A63"
COLOR_MEDIDOR = "E29B2E"
COLOR_EXITO = "2E7D5B"
COLOR_PELIGRO = "C1432B"
COLOR_FONDO = "F7F8F6"
COLOR_BLANCO = "FFFFFF"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, hex_color="CCCCCC"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

DEFECTOS = [
    {
        "codigo": "DEF-2026-001",
        "archivo": "Vulnerabilidad_Sesion_JWT_Expirada_Y_Redireccion",
        "titulo_corto": "Módulo: Seguridad — Manejo de Token Expirado y Redirección Insegura",
        "modulo": "Seguridad & Autenticación API",
        "severidad": "Crítica (Fallo en autenticación y posibles bloqueos de sesión)",
        "prioridad": "P1 (Alta)",
        "tipo": "Seguridad / Manejo de Sesión",
        "descripcion": (
            "Al expirar el token JWT de sesión (IDX10223: Lifetime validation failed), el "
            "middleware de redirección HTTPS fallaba al determinar el puerto de redirección "
            "(HttpsRedirectionMiddleware Failed to determine the https port), provocando "
            "rechazos silenciosos o errores 500 que impedían el refresco controlado y "
            "cerraban abruptamente la sesión del usuario en FuelTrack."
        ),
        "pasos": [
            "1. Iniciar sesión con un usuario válido en el frontend de FuelTrack.",
            "2. Permitir que transcurra el tiempo de expiración del token JWT configurado.",
            "3. Ejecutar cualquier acción que consuma un endpoint protegido (ej: consultar tanques o recepciones).",
            "4. Observar la excepción SecurityTokenExpiredException y la falta de redirección limpia a login."
        ],
        "observado": "- Excepción no interceptada con error 500 en consola de API y pantalla congelada en cliente.",
        "esperado": "- Intercepción controlada de código 401/403 con cierre de sesión seguro y redirección inmediata a /login.",
        "rca": (
            "Manejo deficiente de la revocación automática de tokens expirados en el pipeline de "
            "autenticación y falta de captura específica en el cliente HTTP de frontend para "
            "invalidar la sesión sin fugas de estado."
        ),
        "solucion": [
            ("Backend (FuelTrack.Api/Program.cs):", "Configuración de manejo explícito de errores de autenticación JWT."),
            ("Frontend (src/services/api.js):", "Captura de error 401 por expiración, borrado de token en localStorage y redirección ordenada a login.")
        ],
        "verificacion": [
            ("Pruebas Unitarias de Seguridad", "dotnet test backend/FuelTrack.slnx", "Aprobado (284/284)"),
            ("Compilación de Producción Frontend", "npm.cmd run build", "Aprobado (64 módulos)")
        ]
    },
    {
        "codigo": "DEF-2026-002",
        "archivo": "Ausencia_Generacion_OnDemand_PDF_Cierre_Diario",
        "titulo_corto": "Módulo: Cierre Diario — Generación y Persistencia On-Demand de Acta PDF",
        "modulo": "Cierre Diario & Auditoría de Operaciones",
        "severidad": "Mayor (Imposibilidad de auditar y exportar actas sin PDF previo)",
        "prioridad": "P1 (Alta)",
        "tipo": "Flujo de Negocio / Auditoría Oficial",
        "descripcion": (
            "Cuando un cierre diario era registrado, si el archivo PDF (PdfActa) no se había "
            "procesado previamente en el background o permanecía nulo, el endpoint "
            "/cierres-diarios/{id}/pdf lanzaba un error 404 - PDF_NO_DISPONIBLE. El frontend "
            "deshabilitaba permanentemente el botón de descarga, bloqueando la obtención del "
            "acta formal de auditoría requerida por la administración."
        ),
        "pasos": [
            "1. Navegar a Cierre Diario y generar un cierre para la fecha actual.",
            "2. Intentar descargar el Acta Oficial en PDF directamente desde la tabla o detalle.",
            "3. Verificar que el botón permanecía deshabilitado o retornaba 404 si PdfActa era null."
        ],
        "observado": "- El botón quedaba inhabilitado permanentemente y el acta no podía ser descargada.",
        "esperado": "- Generación al vuelo (on-demand) del documento PDF formal en caso de no existir previamente, persistencia en BD y descarga inmediata.",
        "rca": (
            "El método CierreDiarioService.GetPdfAsync únicamente leía la columna binaria de BD; "
            "si estaba vacía, no ejecutaba el motor de QuestPDF ni persistía el resultado."
        ),
        "solucion": [
            ("Backend (CierreDiarioService.cs):", "Implementación de generación al vuelo si PdfActa es null, persistencia con SaveChangesAsync() y rediseño completo de plantilla QuestPDF con balance por tanque."),
            ("Frontend (CierreDiarioPage.jsx):", "Habilitación de descarga directa con apiDownload, tarjetas KPI de resumen histórico y modal informativo de balance.")
        ],
        "verificacion": [
            ("Pruebas Backend CierreDiario", "dotnet test backend/FuelTrack.slnx", "Aprobado (284/284)"),
            ("Linter Frontend", "npm.cmd run lint", "Aprobado (0 errores)")
        ]
    },
    {
        "codigo": "DEF-2026-003",
        "archivo": "Falta_Sellos_Criptograficos_QR_Y_Firmas_Ticket_PDF",
        "titulo_corto": "Módulo: Despachos — Sellos Criptográficos, QR y Triple Firma en Tickets PDF",
        "modulo": "Despachos & Tickets de Combustible",
        "severidad": "Mayor (Riesgo de fraude físico y tickets no verificables en estación)",
        "prioridad": "P1 (Alta)",
        "tipo": "Seguridad Física / Integridad Operativa",
        "descripcion": (
            "El formato PDF de los tickets de combustible emitidos carecía de código QR para lectura "
            "rápida en estación de servicio móvil, no incluía el sello criptográfico SHA-256 de "
            "integridad y no poseía las tres líneas reglamentarias de firma (Solicitante, "
            "Operador de Estación y Supervisor)."
        ),
        "pasos": [
            "1. Ir al módulo de Tickets y seleccionar Descargar PDF de un ticket autorizado.",
            "2. Abrir el documento generado.",
            "3. Observar la ausencia de QR de escaneo, hash de seguridad y líneas de responsabilidad."
        ],
        "observado": "- Formato de ticket plano, sin garantías de no-repudio ni soporte para lectores de campo.",
        "esperado": "- Documento formal con código QR embebido, sello criptográfico HMAC-SHA256 y tres casillas de firmas reglamentarias.",
        "rca": (
            "Plantilla inicial básica de QuestPDF que solo imprimía datos crudos de texto sin "
            "estructurar el documento conforme a los requisitos de control fiscal y seguridad de FuelTrack."
        ),
        "solucion": [
            ("Backend (TicketPdfService.cs):", "Reescritura completa con QuestPDF, generación dinámica de código QR con QRCoder, huella SHA-256 y bloque de firmas con tokens corporativos (#16333A, #4A5A63)."),
            ("Frontend (TicketsPage.jsx):", "Descarga directa de PDF desde tabla y modal, con vista previa interactiva del código.")
        ],
        "verificacion": [
            ("Pruebas Unitarias de Tickets", "dotnet test backend/FuelTrack.slnx", "Aprobado (284/284)"),
            ("Compilación Frontend", "npm.cmd run build", "Aprobado")
        ]
    },
    {
        "codigo": "DEF-2026-004",
        "archivo": "Inexistencia_Validaciones_Dominio_Republica_Dominicana",
        "titulo_corto": "Módulo: Formularios — Validaciones Específicas Dominicanas (RNC, Cédula, Teléfono)",
        "modulo": "Validación de Entradas & Catálogos",
        "severidad": "Crítica (Contaminación de base de datos con registros fiscales y personas erróneas)",
        "prioridad": "P1 (Alta)",
        "tipo": "Integridad de Datos / Reglas de Negocio",
        "descripcion": (
            "Los formularios del sistema permitían guardar RNCs con cualquier número de caracteres, "
            "cédulas inválidas sin los 11 dígitos requeridos en República Dominicana, números de teléfono "
            "con prefijos no correspondientes al país (solo se permiten 809, 829 y 849) y placas "
            "vehiculares con formatos alfanuméricos incorrectos."
        ),
        "pasos": [
            "1. Abrir modal de Nuevo Proveedor e ingresar un RNC con 4 caracteres numéricos.",
            "2. Abrir modal de Nuevo Empleado e ingresar un teléfono con código de área 555.",
            "3. Observar que el formulario permitía el envío sin validar el estándar nacional de RD."
        ],
        "observado": "- Envío permitido de identificadores inválidos que generaban fallos posteriores o ensuciaban el catálogo.",
        "esperado": "- Validación estricta en tiempo real con mensajes explicativos y formateo visual conforme al marco regulatorio dominicano.",
        "rca": (
            "Inexistencia de un módulo centralizado de validación con las reglas de negocio de República "
            "Dominicana y dependencia de inputs HTML estándar sin validadores asociados."
        ),
        "solucion": [
            ("Frontend (src/utils/validators.js):", "Creación de funciones validateRnc, validateCedula, validateTelefonoRD (809/829/849), validatePlacaVehiculo, validateAnioVehiculo y formatters asociados."),
            ("Frontend (Field.jsx & Páginas):", "Conexión de errores en tiempo real y componentes visuales en Proveedores, Empleados y Vehículos.")
        ],
        "verificacion": [
            ("Linter Frontend", "npm.cmd run lint", "Aprobado (0 errores)"),
            ("Compilación de Producción", "npm.cmd run build", "Aprobado")
        ]
    },
    {
        "codigo": "DEF-2026-005",
        "archivo": "Deficiencia_Contraste_StatusBadge_Y_Ergonomia_Tactil",
        "titulo_corto": "Módulo: UI/UX — Contraste de Badges y Botones de Acción Ergonómicos (38px)",
        "modulo": "UI/UX & Factores Humanos",
        "severidad": "Media (Violación de normas de accesibilidad y fatiga operativa en dispositivos)",
        "prioridad": "P2 (Media)",
        "tipo": "Accesibilidad & Usabilidad",
        "descripcion": (
            "Las insignias de estado (StatusBadge) en tono amarillo presentaban texto blanco o claro "
            "sobre fondo brillante, violando los criterios de contraste WCAG AA. Además, los botones de "
            "acción en las tablas medían menos de 28px de altura, dificultando la pulsación en tablets o "
            "pantallas táctiles de operadores de combustible."
        ),
        "pasos": [
            "1. Inspeccionar las insignias de estado 'Pendiente' o 'En Proceso' en la tabla de solicitudes.",
            "2. Medir el ratio de contraste tipográfico contra el fondo.",
            "3. Intentar hacer clic en los botones de acción en una pantalla de baja resolución o táctil."
        ],
        "observado": "- Texto ilegible en badges amarillos y botones minúsculos con riesgo de clics erróneos.",
        "esperado": "- Badges con fondos suaves translúcidos y texto oscurecido de alto contraste (#854d0e); botones de al menos 38px con íconos vectoriales claros.",
        "rca": (
            "Uso de clases de Tailwind con colores saturados directos (bg-yellow-500 text-white) en vez "
            "de tokens corporativos calibrados de accesibilidad."
        ),
        "solucion": [
            ("Frontend (StatusBadge.jsx):", "Rediseño completo con paleta de alto contraste basada en tokens institucionales (#854d0e sobre bg-medidor/15)."),
            ("Frontend (Tablas y Páginas):", "Estandarización de botones con min-height de 38px, íconos SVG y bordes sutiles.")
        ],
        "verificacion": [
            ("Linter Frontend", "npm.cmd run lint", "Aprobado"),
            ("Compilación de Producción", "npm.cmd run build", "Aprobado")
        ]
    },
    {
        "codigo": "DEF-2026-006",
        "archivo": "Falta_Validacion_Espacio_Tanque_Y_Comprobante_Recepciones",
        "titulo_corto": "Módulo: Recepciones — Alerta de Capacidad en Tanque y Comprobante Digital",
        "modulo": "Recepciones de Combustible",
        "severidad": "Mayor (Riesgo de sobrellenado físico en tanques y falta de comprobante)",
        "prioridad": "P1 (Alta)",
        "tipo": "Seguridad Operacional / Control de Inventario",
        "descripcion": (
            "El formulario de recepciones de combustible permitía ingresar cualquier volumen recibido "
            "sin verificar si el tanque de destino tenía espacio disponible (Capacidad - NivelActual). "
            "Adicionalmente, el detalle de recepción se mostraba en una lista plana de textos sin opción "
            "de impresión formal."
        ),
        "pasos": [
            "1. Seleccionar un tanque con 500 galones disponibles.",
            "2. Registrar una recepción de 2,000 galones en dicho tanque.",
            "3. Observar que el formulario no advertía del desbordamiento y registraba la transacción."
        ],
        "observado": "- Registro ciego de volúmenes superiores a la capacidad física del tanque receptor.",
        "esperado": "- Verificación dinámica de espacio libre con advertencia en rojo y modal estructurado como Comprobante Digital con botón de impresión.",
        "rca": (
            "Falta de cálculo comparativo en cliente entre la capacidad del tanque seleccionado y el "
            "input numérico de volumen ingresado por el usuario."
        ),
        "solucion": [
            ("Frontend (RecepcionesPage.jsx):", "Cálculo en vivo de espacio disponible, advertencia de sobrellenado, validación de factura/conduce y rediseño de modal como Comprobante Digital de Recepción imprimible con window.print().")
        ],
        "verificacion": [
            ("Compilación Frontend", "npm.cmd run build", "Aprobado"),
            ("Linter Frontend", "npm.cmd run lint", "Aprobado")
        ]
    },
    {
        "codigo": "DEF-2026-007",
        "archivo": "Ambiguedad_Nomenclatura_Tanques_Y_Nivel_Critico",
        "titulo_corto": "Módulo: Tanques — Asistente de Nomenclatura y Regla de Nivel Crítico Menor",
        "modulo": "Tanques Estacionarios & Almacenamiento",
        "severidad": "Media (Inconsistencia en identificadores y alertas de inventario desactivadas)",
        "prioridad": "P2 (Media)",
        "tipo": "Integridad de Datos / Reglas de Negocio",
        "descripcion": (
            "Se permitía registrar tanques con códigos arbitrarios sin estándar unificado y con un "
            "nivel crítico igual o mayor que la capacidad máxima del tanque, lo que provocaba que el "
            "sistema generara alertas falsas permanentes de combustible crítico."
        ),
        "pasos": [
            "1. En Tanques, presionar '+ Nuevo tanque'.",
            "2. Ingresar capacidad de 1,000 galones y nivel crítico de 1,500 galones.",
            "3. Enviar el formulario y verificar la ausencia de bloqueo."
        ],
        "observado": "- Tanque guardado con nivel crítico absurdo superior a la capacidad total.",
        "esperado": "- Bloqueo inmediato del formulario indicando que el nivel crítico debe ser estrictamente menor a la capacidad total, y botón de autogeneración de código.",
        "rca": (
            "Validación aislada de campos numéricos sin regla relacional entre nivel crítico y capacidad."
        ),
        "solucion": [
            ("Frontend (TanquesPage.jsx & validators.js):", "Implementación de validateNivelCriticoTanque(critico, capacidad), botón 'Autogenerar' para códigos TNQ-[TIPO]-[01] y barra visual de capacidad.")
        ],
        "verificacion": [
            ("Compilación Frontend", "npm.cmd run build", "Aprobado"),
            ("Linter Frontend", "npm.cmd run lint", "Aprobado")
        ]
    },
    {
        "codigo": "DEF-2026-008",
        "archivo": "Carencia_Control_Fechas_Invertidas_Reportes_Notificaciones",
        "titulo_corto": "Módulo: Reportes — Bloqueo de Fechas Invertidas (Desde > Hasta)",
        "modulo": "Reportes & Notificaciones",
        "severidad": "Media (Consultas SQL erróneas y exportación de archivos vacíos)",
        "prioridad": "P2 (Media)",
        "tipo": "Validación de Entradas & UX",
        "descripcion": (
            "En las pantallas de Reportes y Notificaciones, el usuario podía establecer una fecha "
            "'Desde' posterior a la fecha 'Hasta'. Al ejecutar la búsqueda o exportar a Excel/PDF, el "
            "backend procesaba una consulta ilógica y retornaba reportes vacíos sin mensaje de ayuda."
        ),
        "pasos": [
            "1. En Reportes, colocar Desde: 2026-10-01 y Hasta: 2026-09-01.",
            "2. Pulsar 'Filtrar' o 'Exportar Excel'.",
            "3. Observar la ejecución de la consulta sin validación de coherencia temporal."
        ],
        "observado": "- Consulta ejecutada y archivo vacío descargado sin indicación del error en fechas.",
        "esperado": "- Restricción mutua con atributos max/min en los selectores de fecha y bloqueo con banner de advertencia si Desde > Hasta.",
        "rca": (
            "Ausencia de enlace cruzado entre el estado de fechaDesde y fechaHasta en el formulario de filtros."
        ),
        "solucion": [
            ("Frontend (ReportesPage.jsx & NotificacionesPage.jsx):", "Integración de validateRangoFechas(), restricciones cruzadas max/min y botones ergonómicos de exportación con íconos.")
        ],
        "verificacion": [
            ("Compilación Frontend", "npm.cmd run build", "Aprobado"),
            ("Linter Frontend", "npm.cmd run lint", "Aprobado")
        ]
    },
    {
        "codigo": "DEF-2026-009",
        "archivo": "Ausencia_Calculo_Rapido_Aprobacion_Solicitudes",
        "titulo_corto": "Módulo: Solicitudes — Cálculo Rápido de Cuota y Control de Autorización",
        "modulo": "Solicitudes de Combustible",
        "severidad": "Media (Demoras en aprobación y falta de confirmación en rechazos)",
        "prioridad": "P2 (Media)",
        "tipo": "Optimización Operativa / Usabilidad",
        "descripcion": (
            "En la aprobación de solicitudes de combustible, el supervisor debía calcular manualmente "
            "fracciones de galones si deseaba autorizar el 75% o 50% de lo solicitado. Además, los rechazos "
            "se ejecutaban sin diálogo de confirmación, pudiendo descartar solicitudes por error."
        ),
        "pasos": [
            "1. En Solicitudes, abrir la aprobación de una solicitud de 85 galones.",
            "2. Notar la falta de opciones rápidas para autorizar un porcentaje.",
            "3. Notar la falta de confirmación al presionar rechazar."
        ],
        "observado": "- Operación lenta y propensa a pulsaciones no intencionadas de rechazo.",
        "esperado": "- Modal con tarjeta contextual del vehículo/solicitante, botones de cuota rápida (100%, 75%, 50%) y confirmación modal de rechazo.",
        "rca": (
            "Modal de aprobación genérico sin enriquecimiento ergonómico para toma de decisiones rápida."
        ),
        "solucion": [
            ("Frontend (SolicitudesPage.jsx):", "Rediseño completo del modal de aprobación con botones rápidos 100%, 75%, 50%, advertencia si se excede lo solicitado y modal explícito para confirmación de rechazo.")
        ],
        "verificacion": [
            ("Compilación Frontend", "npm.cmd run build", "Aprobado"),
            ("Linter Frontend", "npm.cmd run lint", "Aprobado")
        ]
    },
    {
        "codigo": "DEF-2026-010",
        "archivo": "Falta_Nomenclatura_Controlada_Y_Auditoria_Tickets",
        "titulo_corto": "Módulo: Tickets — Selector de Prefijos Operativos y Desacoplamiento de Tabla",
        "modulo": "Gestión de Tickets de Despacho",
        "severidad": "Mayor (Saturación de acciones en tabla y códigos de ticket no estandarizados)",
        "prioridad": "P1 (Alta)",
        "tipo": "UI/UX & Flujo Operativo",
        "descripcion": (
            "Al crear un ticket, el código debía escribirse a mano sin validación de prefijo oficial; "
            "asimismo, la tabla de tickets mostraba hasta 5 botones en cada fila (Emitir, Enviar, Anular, "
            "Detalle, PDF), provocando una sobrecarga visual severa que dificultaba la lectura en pantalla."
        ),
        "pasos": [
            "1. Crear un ticket y escribir un código arbitrario sin prefijo.",
            "2. Revisar la tabla principal y notar la multiplicidad de botones por fila."
        ],
        "observado": "- Dispersión de acciones críticas en filas de la tabla sin control de prefijo.",
        "esperado": "- Selector visual de prefijos oficiales (COM, TCK, DSL, GAS, EMG) con vista previa correlativa y tabla limpia con acciones consolidadas en el modal de auditoría.",
        "rca": (
            "Diseño de interfaz no optimizado para la densidad de información de un ERP de combustible."
        ),
        "solucion": [
            ("Frontend (TicketsPage.jsx):", "Píldoras selectoras de prefijo con código en vivo (COM-2026-XXXXXX), tabla estilizada con botón directo 'Detalle' y 'Acta PDF', y consolidación de acciones operativas dentro del modal de auditoría.")
        ],
        "verificacion": [
            ("Compilación Frontend", "npm.cmd run build", "Aprobado"),
            ("Linter Frontend", "npm.cmd run lint", "Aprobado")
        ]
    }
]

def create_report_doc(d, output_dir):
    doc = Document()
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("SISTEMA FUELTRACK ERP  |  ASEGURAMIENTO DE CALIDAD (QA)  |  " + d["codigo"])
        hr.font.name = "Arial"
        hr.font.size = Pt(8)
        hr.font.color.rgb = RGBColor(0x4A, 0x5A, 0x63)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = fp.add_run("INFORME TÉCNICO OFICIAL DE DEFECTO Y RESOLUCIÓN  •  CONFIDENCIAL INSTITUCIONAL")
        fr.font.name = "Arial"
        fr.font.size = Pt(8)
        fr.font.color.rgb = RGBColor(0x4A, 0x5A, 0x63)

    # Titulo Principal Banner
    t_banner = doc.add_table(rows=1, cols=1)
    t_banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_banner = t_banner.rows[0].cells[0]
    set_cell_background(c_banner, COLOR_TANQUE)
    set_cell_margins(c_banner, top=160, bottom=160, left=200, right=200)
    p_b = c_banner.paragraphs[0]
    p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_b1 = p_b.add_run("INFORME TÉCNICO DE NO CONFORMIDAD Y RESOLUCIÓN\n")
    r_b1.bold = True
    r_b1.font.name = "Arial"
    r_b1.font.size = Pt(14)
    r_b1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    r_b2 = p_b.add_run(f"{d['codigo']} — {d['titulo_corto']}")
    r_b2.bold = True
    r_b2.font.name = "Arial"
    r_b2.font.size = Pt(11)
    r_b2.font.color.rgb = RGBColor(0xE2, 0x9B, 0x2E)

    doc.add_paragraph()

    # Tabla de Metadatos
    t_meta = doc.add_table(rows=4, cols=4)
    t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_meta, "B0C4DE")
    
    meta_rows = [
        [("Código:", True), (d["codigo"], False), ("Fecha de Cierre:", True), ("18/09/2026", False)],
        [("Módulo Afectado:", True), (d["modulo"], False), ("Tipo Defecto:", True), (d["tipo"], False)],
        [("Severidad:", True), (d["severidad"], False), ("Prioridad:", True), (d["prioridad"], False)],
        [("Estado QA:", True), ("✅ CERRADO / RESUELTO", False), ("Auditor Líder:", True), ("Rowlis Trinidad", False)]
    ]
    
    for r_idx, row_data in enumerate(meta_rows):
        for c_idx, (text, is_bold) in enumerate(row_data):
            cell = t_meta.cell(r_idx, c_idx)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            if is_bold:
                set_cell_background(cell, COLOR_FONDO)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.size = Pt(9)
            run.bold = is_bold
            if is_bold:
                run.font.color.rgb = RGBColor(0x16, 0x33, 0x3A)
            elif "CERRADO" in text:
                run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x5B)
                run.bold = True
            else:
                run.font.color.rgb = RGBColor(0x12, 0x18, 0x1A)

    doc.add_paragraph()

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x16, 0x33, 0x3A)

    # 1. Descripción
    add_section_header("1. DESCRIPCIÓN Y CONTEXTO DEL DEFECTO")
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.line_spacing = 1.15
    r_desc = p_desc.add_run(d["descripcion"])
    r_desc.font.name = "Arial"
    r_desc.font.size = Pt(9.5)

    # 2. Pasos de Reproducción
    add_section_header("2. PROCEDIMIENTO DE REPRODUCCIÓN (PASO A PASO)")
    for step in d["pasos"]:
        p_step = doc.add_paragraph()
        p_step.paragraph_format.left_indent = Inches(0.2)
        p_step.paragraph_format.space_after = Pt(2)
        r_step = p_step.add_run(step)
        r_step.font.name = "Arial"
        r_step.font.size = Pt(9)

    # 3. Comportamiento Observado vs Esperado
    add_section_header("3. EVALUACIÓN DE COMPORTAMIENTOS")
    t_comp = doc.add_table(rows=2, cols=2)
    t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_comp, "CCCCCC")
    
    h0 = t_comp.cell(0, 0)
    h1 = t_comp.cell(0, 1)
    set_cell_background(h0, "F8D7DA")
    set_cell_background(h1, "D4EDDA")
    set_cell_margins(h0, top=100, bottom=100, left=120, right=120)
    set_cell_margins(h1, top=100, bottom=100, left=120, right=120)
    
    rh0 = h0.paragraphs[0].add_run("COMPORTAMIENTO ANTERIOR (ERRÓNEO)")
    rh0.bold = True
    rh0.font.name = "Arial"
    rh0.font.size = Pt(9)
    rh0.font.color.rgb = RGBColor(0x72, 0x1C, 0x24)

    rh1 = h1.paragraphs[0].add_run("COMPORTAMIENTO ESPERADO (CORRECTO)")
    rh1.bold = True
    rh1.font.name = "Arial"
    rh1.font.size = Pt(9)
    rh1.font.color.rgb = RGBColor(0x15, 0x57, 0x24)

    c0 = t_comp.cell(1, 0)
    c1 = t_comp.cell(1, 1)
    set_cell_margins(c0, top=100, bottom=100, left=120, right=120)
    set_cell_margins(c1, top=100, bottom=100, left=120, right=120)
    
    rc0 = c0.paragraphs[0].add_run(d["observado"])
    rc0.font.name = "Arial"
    rc0.font.size = Pt(9)
    
    rc1 = c1.paragraphs[0].add_run(d["esperado"])
    rc1.font.name = "Arial"
    rc1.font.size = Pt(9)

    # 4. RCA
    add_section_header("4. ANÁLISIS DE CAUSA RAÍZ (RCA - ROOT CAUSE ANALYSIS)")
    p_rca = doc.add_paragraph()
    p_rca.paragraph_format.line_spacing = 1.15
    r_rca = p_rca.add_run(d["rca"])
    r_rca.font.name = "Arial"
    r_rca.font.size = Pt(9.5)

    # 5. Solución
    add_section_header("5. RESOLUCIÓN TÉCNICA E IMPLEMENTACIÓN")
    for comp, sol_text in d["solucion"]:
        p_sol = doc.add_paragraph()
        p_sol.paragraph_format.left_indent = Inches(0.15)
        p_sol.paragraph_format.space_after = Pt(3)
        r_c = p_sol.add_run(comp + " ")
        r_c.bold = True
        r_c.font.name = "Arial"
        r_c.font.size = Pt(9)
        r_c.font.color.rgb = RGBColor(0x16, 0x33, 0x3A)
        r_t = p_sol.add_run(sol_text)
        r_t.font.name = "Arial"
        r_t.font.size = Pt(9)

    # 6. Verificación
    add_section_header("6. MATRIZ DE PRUEBAS Y VALIDACIÓN AUTOMATIZADA")
    t_ver = doc.add_table(rows=len(d["verificacion"]) + 1, cols=3)
    t_ver.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_ver, "CCCCCC")
    
    headers_ver = ["Tipo de Prueba", "Comando / Suite Ejecutada", "Resultado Oficial"]
    for i, h in enumerate(headers_ver):
        c = t_ver.cell(0, i)
        set_cell_background(c, COLOR_ACERO)
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        rh = c.paragraphs[0].add_run(h)
        rh.bold = True
        rh.font.name = "Arial"
        rh.font.size = Pt(8.5)
        rh.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, (t_name, t_cmd, t_res) in enumerate(d["verificacion"], start=1):
        c0 = t_ver.cell(r_idx, 0)
        c1 = t_ver.cell(r_idx, 1)
        c2 = t_ver.cell(r_idx, 2)
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c2, top=60, bottom=60, left=100, right=100)
        
        c0.paragraphs[0].add_run(t_name).font.size = Pt(8.5)
        c1.paragraphs[0].add_run(t_cmd).font.size = Pt(8.5)
        r_res = c2.paragraphs[0].add_run(t_res)
        r_res.font.size = Pt(8.5)
        r_res.bold = True
        r_res.font.color.rgb = RGBColor(0x2E, 0x7D, 0x5B)

    # 7. Firmas
    add_section_header("7. CERTIFICACIÓN DE CALIDAD Y APROBACIÓN INSTITUCIONAL")
    t_firmas = doc.add_table(rows=1, cols=3)
    t_firmas.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_firmas, "FFFFFF")
    
    firmas_data = [
        ("Rowlis Trinidad", "Ingeniero Líder de QA", "Aprobado & Certificado"),
        ("Equipo de Desarrollo", "Arquitectura de Software", "Implementación Completa"),
        ("Dirección de Operaciones", "Revisión y Conformidad", "Visto Bueno Producción")
    ]
    
    for i, (f_nom, f_cargo, f_est) in enumerate(firmas_data):
        c = t_firmas.cell(0, i)
        set_cell_margins(c, top=140, bottom=60, left=60, right=60)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        rf0 = p.add_run("_______________________________\n")
        rf0.font.color.rgb = RGBColor(0x4A, 0x5A, 0x63)
        
        rf1 = p.add_run(f_nom + "\n")
        rf1.bold = True
        rf1.font.name = "Arial"
        rf1.font.size = Pt(8.5)
        rf1.font.color.rgb = RGBColor(0x16, 0x33, 0x3A)
        
        rf2 = p.add_run(f_cargo + "\n")
        rf2.font.name = "Arial"
        rf2.font.size = Pt(8)
        rf2.font.color.rgb = RGBColor(0x4A, 0x5A, 0x63)
        
        rf3 = p.add_run(f_est)
        rf3.font.name = "Arial"
        rf3.font.size = Pt(7.5)
        rf3.font.color.rgb = RGBColor(0x2E, 0x7D, 0x5B)

    out_file = os.path.join(output_dir, f"{d['codigo']}_{d['archivo']}.docx")
    doc.save(out_file)
    print(f"[OK] Generado: {d['codigo']}_{d['archivo']}.docx")

def main():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    print(f"Generando {len(DEFECTOS)} informes técnicos individuales en .docx...")
    for d in DEFECTOS:
        create_report_doc(d, base_dir)
    print("Todos los informes individuales han sido generados exitosamente!")

if __name__ == "__main__":
    main()
